"""Export to Canva: editable structure, geometry, fidelity, pages and the Canva API.

What these tests can and cannot show: the PowerPoint files are checked by
reading them back (python-pptx) and by rendering them with an independent
program (LibreOffice) against the original presentation (Chromium). The
Canva Connect protocol is checked against ``tests/canva_mock.py``. Canva's
own import and rendering cannot run here (canva.com is not reachable from
this environment, and no Canva account is available).
"""

from __future__ import annotations

import shutil
import socket
import zipfile

import numpy as np
import pytest
from PIL import Image, ImageDraw
from pptx import Presentation as Pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from canva_mock import MockCanva
from behance_studio.canva.build import BuildOptions, build_scene
from behance_studio.canva.check import compare, libreoffice
from behance_studio.canva.connect import CanvaConnect, CanvaError
from behance_studio.canva.export import CanvaOptions, export_to_canva
from behance_studio.canva.model import Gradient
from behance_studio.canva.pages import AUTO_RANGE, paginate
from behance_studio.canva.pptx_out import emu, write_pptx
from behance_studio.layout.model import Block, Presentation, Section, TextStyle
from behance_studio.render.browser import browser_available

HAS_BROWSER = browser_available()


@pytest.fixture(scope="module")
def generated(tmp_path_factory, demo_assets_dir):
    """A project with one proposal per layout family (no previews)."""
    from behance_studio.assets.importer import import_paths
    from behance_studio.core import settings
    from behance_studio.generation.pipeline import GenerationRequest, generate
    from behance_studio.layout.families import available_families
    from behance_studio.projects.store import Project

    home = tmp_path_factory.mktemp("canva_home")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("BEHANCE_STUDIO_HOME", str(home))
        settings.reset_settings_store_for_tests()
        project = Project.create(tmp_path_factory.mktemp("canva_projects"), "Butter & Bean")
        import_paths(project, [demo_assets_dir])
        generate(project, GenerationRequest(families=available_families(), seed=4, render_previews=False))
    settings.reset_settings_store_for_tests()
    return project


def _proposals(project, families=None):
    from behance_studio.projects.versions import VersionManager

    vm = VersionManager(project)
    return [(r, vm.load(r.id).presentation) for r in vm.list() if families is None or r.family in families]


@pytest.fixture
def no_keyring(monkeypatch):
    from behance_studio.ai import credentials

    monkeypatch.setattr(credentials, "_keyring", lambda: None)
    credentials._session.clear()
    CanvaConnect._access.clear()
    yield
    credentials._session.clear()
    CanvaConnect._access.clear()


@pytest.fixture
def mock():
    server = MockCanva()
    yield server
    server.close()


def _free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


# ------------------------------------------------------------- structure
def test_every_layer_becomes_its_own_editable_element(generated, tmp_path):
    for record, pres in _proposals(generated):
        scene = build_scene(generated, pres, BuildOptions(), tmp_path / record.family)
        by_block: dict[str, list] = {}
        for e in scene.elements:
            by_block.setdefault(e.block_id, []).append(e)
            # nothing is a flattened page: every picture is one layer or a section background
            if e.kind == "image":
                assert e.block_id or e.role == "background", (record.family, e.name)
        y = 0.0
        for index, section in enumerate(pres.sections):
            bgs = [e for e in scene.elements if e.role == "background" and e.section == index]
            assert len(bgs) == 1 and bgs[0].w == pres.width and abs(bgs[0].y - y) < 0.01 and abs(bgs[0].h - section.height) < 0.01
            for b in section.blocks:
                els = by_block.get(b.id, [])
                if b.kind == "text" and any(line.strip() for line in b.lines):
                    texts = [e for e in els if e.kind == "text"]
                    assert len(texts) == 1, (record.family, b.role)
                    assert texts[0].text.lines == b.lines and texts[0].text.family == b.style.family
                    assert abs(texts[0].text.size - b.style.size) < 1e-6 and texts[0].text.line_px == b.style.line_px
                elif b.kind in ("image", "video") and b.asset_id:
                    media = [e for e in els if e.kind == "image" and e.picture.asset_id == b.asset_id]
                    assert media, (record.family, b.kind, b.asset_id)
                elif b.kind == "shape" and (b.fill or b.stroke or b.fill_asset):
                    assert els, (record.family, b.shape)
            y += section.height
        assert scene.fonts, record.family


def test_pages_have_one_size_and_never_cut_text(generated, tmp_path):
    for record, pres in _proposals(generated):
        scene = build_scene(generated, pres, BuildOptions(), tmp_path / record.family)
        paged = paginate(scene)
        H = paged.page_height
        assert AUTO_RANGE[0] <= H <= AUTO_RANGE[1] or H == round(scene.height)
        assert len(paged.pages) <= 100
        assert paged.pages[0].y0 == 0 and abs(paged.pages[-1].y1 - scene.height) < 0.01
        for a, b in zip(paged.pages, paged.pages[1:]):
            assert a.y1 == b.y0  # consecutive strips: nothing lost or repeated
        def key(e):
            return (e.section, e.block_id, e.role, e.kind, round(e.x, 2), round(e.w, 2), e.picture.path if e.picture else "")

        shown: dict[tuple, float] = {}
        for page in paged.pages:
            assert abs(page.filler - (H - (page.y1 - page.y0))) < 0.01
            for e in page.elements:
                _x0, y0, _x1, y1 = e.bounds()
                if e.kind == "text":
                    assert y0 >= -0.5 and y1 <= H + 0.5, (record.family, page.index, e.name)  # never cut
                if e.splittable and e.role != "filler":
                    shown[key(e)] = shown.get(key(e), 0.0) + e.h
        for e in scene.elements:
            if e.splittable and key(e) in shown:
                assert abs(shown[key(e)] - e.h) < 0.6, (record.family, e.name, shown[key(e)], e.h)  # split pieces add up


def _shape_kind(shape) -> str:
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "image"
    if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
        return "text"
    return "shape"


def test_pptx_is_native_and_matches_the_scene(generated, tmp_path):
    for record, pres in _proposals(generated, {"minimal", "layered", "story", "posters"}):
        scene = build_scene(generated, pres, BuildOptions(), tmp_path / "m" / record.family)
        paged = paginate(scene)
        path = write_pptx(paged, tmp_path / f"{record.family}.pptx")
        deck = Pptx(str(path))
        assert deck.slide_width == emu(pres.width) and deck.slide_height == emu(paged.page_height)
        assert len(deck.slides) == len(paged.pages)
        for slide, page in zip(deck.slides, paged.pages):
            shapes = list(slide.shapes)
            elements = [e for e in page.elements if e.w > 0.01 and e.h > 0.01]
            assert len(shapes) == len(elements)
            kinds = {_shape_kind(s) for s in shapes}
            assert len(shapes) >= 2  # never one flattened picture per page
            for shape, e in zip(shapes, elements):  # same stacking order
                want = "text" if e.kind == "text" else "image" if e.kind == "image" else "shape"
                assert _shape_kind(shape) == want, (record.family, e.name)
                assert abs(shape.left - emu(e.x)) <= 1 and abs(shape.top - emu(e.y)) <= 1
                assert abs(shape.width - max(1, emu(e.w))) <= 1 and abs(shape.height - max(1, emu(e.h))) <= 1
                assert abs((shape.rotation - e.rotation % 360 + 180) % 360 - 180) < 0.01
                assert shape._element.find(qn("p:style")) is None  # no theme outline/shadow
                if e.kind == "text":
                    t = e.text
                    runs = [r for p in shape.text_frame.paragraphs for r in p.runs]
                    assert runs and all(r.font.name == t.family for r in runs)
                    assert all(abs(r.font.size.pt - round(t.size * 0.75 * 100) / 100) < 0.01 for r in runs)
                    assert all(bool(r.font.bold) == (t.weight >= 600) for r in runs)
                    if not t.stroke:
                        assert all(str(r.font.color.rgb) == t.color.lstrip("#").upper() for r in runs)
                    text = shape.text_frame.text.replace("\x0b", "\n")
                    joined = [" ".join(x for x in t.lines if x)] if t.single_line else [x for x in t.lines if x]
                    if t.natural:
                        assert text.split() == " ".join(joined).split()
                    else:
                        assert [x for x in text.split("\n") if x] == joined
                    wrap = shape._element.find(qn("p:txBody")).find(qn("a:bodyPr")).get("wrap")
                    assert wrap == ("square" if t.natural else "none")
                if e.kind == "image":
                    l, tp, r, b = e.picture.crop
                    assert abs(shape.crop_left - l) < 1e-4 and abs(shape.crop_top - tp) < 1e-4
                    assert abs(shape.crop_right - r) < 1e-4 and abs(shape.crop_bottom - b) < 1e-4
            assert "text" in kinds or page.index == len(paged.pages) - 1 or "image" in kinds


# ---------------------------------------------------------------- effects
def _effects_project(projects_dir, tmp_path):
    from behance_studio.assets.importer import import_paths
    from behance_studio.projects.store import Project

    folder = tmp_path / "fx"
    folder.mkdir()
    photo = Image.new("RGB", (1200, 800))
    arr = np.zeros((800, 1200, 3), dtype=np.uint8)
    arr[..., 0] = np.linspace(40, 220, 1200, dtype=np.uint8)[None, :]
    arr[..., 1] = np.linspace(200, 60, 800, dtype=np.uint8)[:, None]
    arr[..., 2] = 120
    photo = Image.fromarray(arr, "RGB")
    ImageDraw.Draw(photo).ellipse((400, 200, 800, 600), fill=(250, 240, 230))
    photo.save(folder / "photo.jpg", quality=92)
    logo = Image.new("RGBA", (600, 600), (255, 255, 255, 0))
    ImageDraw.Draw(logo).ellipse((100, 100, 500, 500), fill=(30, 20, 15, 255))
    logo.save(folder / "logo.png")
    project = Project.create(projects_dir, "Effects")
    import_paths(project, [folder])
    ids = {a.filename: a.id for a in project.list_assets()}
    return project, ids


def _effects_presentation(ids) -> Presentation:
    st = TextStyle("Manrope", 22, 400, 1.4, 0.0, "#202020")
    disp = TextStyle("Fraunces", 64, 700, 1.0, 0.02, "#1A1A1A", uppercase=True)
    b = []
    b.append(Block("image", 40, 40, 400, 260, z=1, asset_id=ids["photo.jpg"], fade="bottom"))
    b.append(Block("image", 480, 40, 400, 260, z=1, asset_id=ids["photo.jpg"], blend="multiply"))
    b.append(Block("image", 920, 40, 400, 260, z=1, asset_id=ids["photo.jpg"], filter="grayscale"))
    b.append(Block("image", 40, 360, 360, 360, z=1, asset_id=ids["photo.jpg"], clip="circle"))
    b.append(Block("image", 440, 360, 300, 420, z=1, asset_id=ids["photo.jpg"], clip="arch"))
    b.append(Block("image", 780, 360, 400, 260, z=1, asset_id=ids["photo.jpg"], perspective=(8.0, -14.0), shadow="soft"))
    b.append(Block("image", 1200, 360, 160, 160, z=1, asset_id=ids["logo.png"], fit="contain", tint="#C2410C"))
    b.append(Block("image", 60, 820, 240, 480, z=1, asset_id=ids["photo.jpg"], device="phone"))
    b.append(Block("image", 340, 820, 640, 420, z=1, asset_id=ids["photo.jpg"], device="browser"))
    b.append(Block("image", 1020, 820, 320, 420, z=1, asset_id=ids["logo.png"], device="frame", fit="contain", mat="#FFFFFF"))
    b.append(Block("shape", 40, 1340, 600, 200, z=1, fill="linear-gradient(to top, rgba(34,24,18,0.78), rgba(34,24,18,0))"))
    b.append(Block("shape", 680, 1340, 600, 200, z=1, fill="radial-gradient(ellipse 70% 60% at 50% 38%, #5F5148 0%, #43342B 72%)", radius=24))
    b.append(Block("shape", 40, 1580, 160, 160, z=1, shape="ring", stroke="#C2410C", stroke_width=6))
    b.append(Block("shape", 240, 1580, 300, 160, z=1, fill="#F4C95D", rotate=-6, shadow="hard"))
    b.append(Block("image", 600, 1580, 200, 200, z=1, asset_id=ids["logo.png"], fit="contain", shadow="lifted"))
    b.append(Block("text", 860, 1580, 480, 64, z=2, text="OUTLINE", lines=["OUTLINE"], style=disp.with_(stroke="#1A1A1A"), role="display"))
    b.append(Block("text", 860, 1680, 420, 64, z=2, text="Tilted", lines=["TILTED"], style=disp, role="display", rotate=-8))
    body = "Butter & Bean is a neighbourhood specialty coffee shop and bakery. The identity brings together carefully sourced coffee."
    from behance_studio.layout.text import wrap

    lines = wrap(body, st, 520)
    b.append(Block("text", 40, 1800, 520, len(lines) * st.line_px, z=2, text=body, lines=lines, style=st, role="body"))
    b.append(Block("text", 0, 1980, 1400, 40, z=2, text="Marquee", lines=["SPECIALTY COFFEE ✦", "FRESHLY BAKED ✦"], style=st.with_(size=26), role="marquee"))
    b.append(Block("text", 640, 1800, 300, 40, z=2, text="Multiply", lines=["Multiply"], style=st.with_(color="#C2410C"), blend="multiply", role="label"))
    b.append(Block("image", 1000, 1800, 200, 150, z=1, asset_id="missing-asset"))
    sections = [Section("custom", 2060, "#F3EEE6", b, title="Effects")]
    return Presentation("custom", 1, 1400, sections, "Effects", fonts=["Manrope", "Fraunces"])


def test_effects_are_reproduced_on_their_own_element(projects_dir, tmp_path):
    project, ids = _effects_project(projects_dir, tmp_path)
    pres = _effects_presentation(ids)
    blocks = pres.sections[0].blocks
    for mode in ("exact", "editable"):
        scene = build_scene(project, pres, BuildOptions(mode=mode), tmp_path / mode)
        by = {}
        for e in scene.elements:
            by.setdefault(e.block_id, []).append(e)
        fade = by[blocks[0].id][0]
        with Image.open(fade.picture.path) as im:
            a = np.asarray(im.convert("RGBA"))[..., 3]
        assert a[2].mean() > 250 and a[-2].mean() < 10 and "bottom edge fade" in fade.picture.baked
        mult = by[blocks[1].id][0]
        with Image.open(mult.picture.path) as im:
            px = np.asarray(im.convert("RGB"), dtype=np.float32)[im.height // 2, 5]
        assert "multiply blend" in mult.picture.baked and px[0] < 0.97 * 255  # darkened by the backdrop
        gray = by[blocks[2].id][0]
        with Image.open(gray.picture.path) as im:
            g = np.asarray(im.convert("RGB"), dtype=np.int16)
        assert np.abs(g[..., 0] - g[..., 1]).max() <= 2 and np.abs(g[..., 1] - g[..., 2]).max() <= 2
        circle = by[blocks[3].id][0]
        arch = by[blocks[4].id][0]
        if mode == "exact":
            with Image.open(circle.picture.path) as im:
                assert im.getpixel((0, 0))[3] == 0 and im.getpixel((im.width // 2, im.height // 2))[3] == 255
            assert circle.geometry == "rect" and "ellipse mask" in circle.picture.baked
        else:
            assert circle.geometry == "ellipse" and not circle.picture.baked
            assert arch.geometry == "round2SameRect" and arch.radius == pytest.approx(150)
        tilt = by[blocks[5].id]
        main = [e for e in tilt if e.kind == "image"][-1]
        assert "perspective tilt" in main.picture.baked and main.rotation == 0 and main.shadow is not None
        tint = by[blocks[6].id][0]
        with Image.open(tint.picture.path) as im:
            r, g_, b_, a_ = im.convert("RGBA").getpixel((im.width // 2, im.height // 2))
        assert (r, g_, b_) == (0xC2, 0x41, 0x0C) and a_ == 255
        phone = by[blocks[7].id]
        assert {e.role for e in phone} >= {"frame", "media"} and len([e for e in phone if e.role == "frame"]) == 2  # body + island
        browser = by[blocks[8].id]
        assert len([e for e in browser if e.kind == "ellipse"]) == 3  # window buttons
        lin, rad = by[blocks[10].id][0], by[blocks[11].id][0]
        if mode == "exact":
            assert lin.kind == "image" and rad.kind == "image"
        else:
            assert isinstance(lin.fill, Gradient) and lin.fill.angle == 0 and rad.kind == "image"  # radial is always drawn
        ring = by[blocks[12].id][0]
        assert ring.kind == "ellipse" and ring.fill is None and ring.line.width == 6
        hard = by[blocks[13].id][0]
        assert hard.rotation == -6 and hard.shadow is not None and hard.shadow.blur == 0
        shadowed = by[blocks[14].id]
        assert [e.role for e in shadowed][0] == "shadow" and shadowed[-1].shadow is None  # transparent box: separate shadow
        outline = by[blocks[15].id][0]
        assert outline.text.stroke == "#1A1A1A"
        assert by[blocks[16].id][0].rotation == -8
        body = by[blocks[17].id][0]
        assert body.text.natural  # body copy re-wraps when edited
        assert by[blocks[18].id][0].text.single_line
        assert by[blocks[19].id][0].text.color != "#C2410C"  # multiply against the background
        missing = by[blocks[20].id][0]
        assert missing.kind == "rect" and "missing" in missing.notes[0]
        paged = paginate(scene)
        path = write_pptx(paged, tmp_path / f"fx-{mode}.pptx")
        xml = zipfile.ZipFile(path).read("ppt/slides/slide1.xml").decode()
        import re

        assert re.search(r'<a:rPr [^>]*><a:ln w="\d+"><a:solidFill><a:srgbClr val="1A1A1A"/></a:solidFill></a:ln><a:noFill/>', xml)  # outlined type
        assert 'wrap="square"' in xml and 'wrap="none"' in xml
        assert "<a:ln w=" in xml and "outerShdw" in xml
        if mode == "editable":
            assert '<a:prstGeom prst="ellipse">' in xml and "gradFill" in xml


@pytest.mark.skipif(not HAS_BROWSER, reason="needs a Chromium-based browser")
def test_effects_conversion_matches_the_original(projects_dir, tmp_path):
    from behance_studio.canva.check import check_fidelity
    from behance_studio.render.browser import BrowserSession

    project, ids = _effects_project(projects_dir, tmp_path)
    pres = _effects_presentation(ids)
    pres.sections[0].blocks = [b for b in pres.sections[0].blocks if b.asset_id != "missing-asset"]
    scene = build_scene(project, pres, BuildOptions(), tmp_path / "w")
    paged = paginate(scene)
    with BrowserSession() as browser:
        result = check_fidelity(project, pres, paged, browser)
    assert result.conversion_similarity >= 0.97, [round(d.similarity, 4) for d in result.conversion]


@pytest.mark.slow
@pytest.mark.skipif(not (HAS_BROWSER and libreoffice() and shutil.which("pdftoppm")), reason="needs Chromium, LibreOffice and pdftoppm")
def test_pptx_rendered_by_libreoffice_matches_the_original(generated, tmp_path):
    """An independent program reading the actual .pptx, compared with the original."""
    from behance_studio.canva.check import check_fidelity
    from behance_studio.render.browser import BrowserSession

    limits = {"minimal": (0.99, 0.98), "layered": (0.99, 0.98), "posters": (0.93, 0.88)}  # posters: weight-850 type shown bold
    with BrowserSession() as browser:
        for record, pres in _proposals(generated, set(limits)):
            scene = build_scene(generated, pres, BuildOptions(), tmp_path / record.family)
            paged = paginate(scene)
            pptx = write_pptx(paged, tmp_path / f"{record.family}.pptx")
            result = check_fidelity(generated, pres, paged, browser, pptx)
            conv, office = limits[record.family]
            assert result.powerpoint is not None and len(result.powerpoint) == len(paged.pages)
            assert result.conversion_similarity >= conv, (record.family, [round(d.similarity, 3) for d in result.conversion])
            assert result.powerpoint_similarity >= office, (record.family, [round(d.similarity, 3) for d in result.powerpoint])


def test_compare_metric_detects_real_differences():
    a = Image.new("RGB", (400, 300), "#F3EEE6")
    b = a.copy()
    ImageDraw.Draw(b).rectangle((50, 50, 250, 200), fill="#1A1A1A")
    assert compare(a, a.copy()).similarity == 1.0
    assert compare(a, b).similarity < 0.8


# ------------------------------------------------------------- Canva API
def test_connect_with_pkce_and_import_a_design(mock, no_keyring, generated, tmp_path):
    from behance_studio.ai import credentials

    record, pres = _proposals(generated, {"minimal"})[0]
    paged = paginate(build_scene(generated, pres, BuildOptions(), tmp_path / "m"))
    pptx = write_pptx(paged, tmp_path / "minimal.pptx")

    client = CanvaConnect(mock.client_id, mock.client_secret, port=_free_port(), api_base=mock.api_base, auth_url=mock.auth_url, timeout=20)
    assert not client.connected
    name = client.authorize(open_url=mock.browser, timeout=20)
    assert name == "Ada Designer" and client.connected
    first_refresh = credentials.get_api_key("canva_refresh_token")
    assert first_refresh and first_refresh.startswith("rt-")

    design = client.import_design(pptx, "Butter & Bean — Minimal Editorial proposal with a long title")
    upload = mock.uploads[-1]
    assert upload["pptx"] and upload["mime_type"].endswith("presentationml.presentation")
    assert upload["title"] == "Butter & Bean — Minimal Editorial proposal with a long title"[:50]
    assert upload["size"] == pptx.stat().st_size
    assert design.edit_url.startswith("https://www.canva.com/design/") and design.edit_url.endswith("/edit")

    # rate limited once: retried after Retry-After
    mock.throttle_imports = 1
    assert client.import_design(pptx, "Again").edit_url
    # expired access token: refreshed once with the single-use refresh token, which is replaced
    mock.expire_access_tokens()
    assert client.import_design(pptx, "After expiry").edit_url
    assert credentials.get_api_key("canva_refresh_token") != first_refresh and mock.refresh[first_refresh] is False
    # a failed conversion is explained
    mock.fail_imports = True
    with pytest.raises(CanvaError) as err:
        client.import_design(pptx, "Broken")
    assert "could not read" in err.value.message and "by hand" in err.value.message
    # disconnect revokes and forgets the sign-in
    refresh = credentials.get_api_key("canva_refresh_token")
    client.disconnect()
    assert not client.connected and refresh in mock.revoked


def test_connect_refusals_are_explained(mock, no_keyring):
    client = CanvaConnect(mock.client_id, mock.client_secret, port=_free_port(), api_base=mock.api_base, auth_url=mock.auth_url, timeout=20)
    with pytest.raises(CanvaError, match="did not allow"):
        client.authorize(open_url=lambda url: mock.browser(url, approve=False), timeout=20)
    wrong = CanvaConnect(mock.client_id, "not-the-secret", port=_free_port(), api_base=mock.api_base, auth_url=mock.auth_url, timeout=20)
    with pytest.raises(CanvaError, match="rejected the sign-in"):
        wrong.authorize(open_url=mock.browser, timeout=20)
    unset = CanvaConnect("", None)
    with pytest.raises(CanvaError, match="not set up"):
        unset.authorize(open_url=lambda url: None)
    url = client.authorization_url("challenge", "state")
    assert "code_challenge_method=s256" in url and "redirect_uri=http%3A%2F%2F127.0.0.1%3A" in url
    assert "design%3Acontent%3Awrite" in url


# ---------------------------------------------------------- export flow
def test_export_to_canva_end_to_end(generated, mock, no_keyring, tmp_path):
    client = CanvaConnect(mock.client_id, mock.client_secret, port=_free_port(), api_base=mock.api_base, auth_url=mock.auth_url, timeout=20)
    client.authorize(open_url=mock.browser, timeout=20)
    records = [r for r, _p in _proposals(generated, {"minimal", "bold"})]
    out = tmp_path / "out"
    report = export_to_canva(generated, [r.id for r in records], CanvaOptions(destination=out, upload=True, check=HAS_BROWSER), client=client)
    assert not report.failures and len(report.designs) == 2
    for d in report.designs:
        assert d.pptx.exists() and d.report.exists() and d.scene_json.exists()
        assert d.edit_url.startswith("https://www.canva.com/design/")
        text = d.report.read_text(encoding="utf-8")
        assert "Already imported into your Canva account" in text and "APPROXIMATIONS" in text and "Page 1:" in text
        if HAS_BROWSER:
            assert d.similarity is not None and d.similarity >= 0.98
        assert d.size_bytes < 50 * 1024 * 1024 and d.pages <= 100
    fonts = report.fonts_dir
    assert fonts is not None and any(fonts.glob("*.ttf")) and any(fonts.glob("*-OFL.txt"))
    exports = [e for e in generated.list_exports() if e.format == "canva"]
    assert len(exports) >= 2 and any(e.details.get("edit_url") for e in exports)
    # all proposals of the project in one design
    combined = export_to_canva(generated, [r.id for r in records], CanvaOptions(destination=out, combine=True, check=False))
    assert len(combined.designs) == 1 and not combined.failures
    one = combined.designs[0]
    deck = Pptx(str(one.pptx))
    assert len(deck.slides) == one.pages >= max(d.pages for d in report.designs) and set(one.version_ids) == {r.id for r in records}
    notes = [s.notes_slide.notes_text_frame.text for s in deck.slides]
    assert any("Minimal" in n for n in notes) and any("Bold" in n for n in notes)  # both proposals, in order


def test_export_center_offers_canva(generated, tmp_path):
    from behance_studio.export.exporter import FORMATS, ExportOptions, export_versions

    assert "canva" in FORMATS
    record = _proposals(generated, {"editorial"})[0][0]
    report = export_versions(generated, [record.id], ExportOptions(formats=["canva"], destination=tmp_path / "ec"))
    assert not report.failures
    files = [p for p in report.files if p.suffix == ".pptx"]
    assert files and len(Pptx(str(files[0])).slides) >= 1


def test_cli_canva_export(generated, tmp_path, capsys):
    from behance_studio import cli

    record = _proposals(generated, {"split"})[0][0]
    code = cli.main(["canva-export", str(generated.root), "--version", str(record.number), "--out", str(tmp_path / "cli"), "--no-check"])
    out = capsys.readouterr().out
    assert code == 0 and ".pptx" in out and "1 Canva design file created" in out
    assert cli.main(["canva-status"]) == 1
    assert "not connected" in capsys.readouterr().out


@pytest.mark.gui
def test_export_dialog_and_settings_panel(qapp, generated, tmp_path, no_keyring):
    from behance_studio.ui.dialogs.canva_export import CanvaExportDialog
    from behance_studio.ui.state import AppState
    from behance_studio.ui.widgets.canva_panel import CanvaPanel
    from test_ui import wait_idle

    state = AppState()
    state.settings.canva.open_after = False
    state.set_project(generated)
    record = _proposals(generated, {"gallery"})[0][0]
    dialog = CanvaExportDialog(state, [record.id])
    assert dialog.checked_ids() == [record.id]
    assert not dialog.upload.isEnabled() and dialog.manual.isChecked()  # not connected: file only
    dialog.dest.setText(str(tmp_path / "dlg"))
    dialog.check.setChecked(False)
    dialog._export()
    wait_idle(state, 300)
    assert dialog.report is not None and len(dialog.report.designs) == 1, dialog.status.text()
    assert dialog.report.designs[0].pptx.exists() and dialog.results.count() >= 2
    dialog.close()

    panel = CanvaPanel(state, lambda m: None)
    panel.client_id.setText("OC-abc")
    panel.secret.setText("secret-xyz")
    panel.save()
    assert state.settings.canva.client_id == "OC-abc" and panel.client().client_secret == "secret-xyz"
    assert panel.redirect.text() == "http://127.0.0.1:51723/oauth/redirect"
    assert "not connected" in panel.status.text().lower() or "set up" in panel.status.text().lower()
