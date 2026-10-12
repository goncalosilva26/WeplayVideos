"""Paged scene → PowerPoint (.pptx) with native, individually editable shapes.

PowerPoint is the format Canva converts most completely into an editable
design (text boxes, pictures and shapes each become their own Canva
element). Every element is written explicitly — position, size, rotation,
fill, outline, shadow, crop, typography — and the theme's default styles
are removed, so nothing depends on PowerPoint defaults.

Units: 1 CSS px = 9525 EMU (96 px per inch), 1 px = 0.75 pt.
"""

from __future__ import annotations

import math
import os
from pathlib import Path
from xml.sax.saxutils import escape

from pptx import Presentation as PptxPresentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu

from behance_studio import APP_NAME, __version__
from behance_studio.canva.model import Element, Gradient, PagedScene, Solid
from behance_studio.core.errors import ExportError

EMU_PER_PX = 9525
PT_PER_PX = 0.75
MAX_SLIDE_EMU = 51206400  # PowerPoint's limit: 56 inches per side

_SHAPES = {"rect": MSO_SHAPE.RECTANGLE, "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE, "ellipse": MSO_SHAPE.OVAL,
           "round2SameRect": MSO_SHAPE.ROUND_2_SAME_RECTANGLE}


def emu(px: float) -> int:
    return int(round(px * EMU_PER_PX))


def _hex(color: str) -> str:
    return (color or "#000000").lstrip("#").upper()[:6].rjust(6, "0")


def _clr(color: str, alpha: float = 1.0) -> str:
    a = "" if alpha >= 0.999 else f'<a:alpha val="{int(round(max(0.0, alpha) * 100000))}"/>'
    return f'<a:srgbClr val="{_hex(color)}">{a}</a:srgbClr>'


def fill_xml(fill, opacity: float = 1.0) -> str:
    if fill is None:
        return "<a:noFill/>"
    if isinstance(fill, Solid):
        return f"<a:solidFill>{_clr(fill.color, fill.alpha * opacity)}</a:solidFill>"
    assert isinstance(fill, Gradient)
    stops = "".join(f'<a:gs pos="{int(round(s.pos * 100000))}">{_clr(s.color, s.alpha * opacity)}</a:gs>' for s in fill.stops)
    if fill.kind == "linear":
        ang = int(round(((fill.angle - 90.0) % 360) * 60000))
        shade = f'<a:lin ang="{ang}" scaled="0"/>'
    else:
        l, t = int(fill.cx * 100000), int(fill.cy * 100000)
        shade = f'<a:path path="circle"><a:fillToRect l="{l}" t="{t}" r="{100000 - l}" b="{100000 - t}"/></a:path>'
    return f'<a:gradFill rotWithShape="1"><a:gsLst>{stops}</a:gsLst>{shade}</a:gradFill>'


def line_xml(el: Element) -> str:
    if el.line is None:
        return "<a:ln><a:noFill/></a:ln>"
    return f'<a:ln w="{emu(el.line.width)}"><a:solidFill>{_clr(el.line.color, el.line.alpha * el.opacity)}</a:solidFill></a:ln>'


def _spread_scale(size: float, spread: float) -> int:
    """CSS shadow spread as a DrawingML shadow scale (1/1000 %)."""
    if not spread or size <= 0:
        return 100000
    return int(round(max(0.05, (size + 2 * spread) / size) * 100000))


def effect_xml(el: Element) -> str:
    s = el.shadow
    if s is None:
        return ""
    alpha = s.alpha * el.opacity
    if s.kind == "glow":
        return f'<a:effectLst><a:glow rad="{emu(s.spread + s.blur * 0.5)}">{_clr(s.color, alpha)}</a:glow></a:effectLst>'
    dist = math.hypot(s.dx, s.dy)
    direction = int(round((math.degrees(math.atan2(s.dy, s.dx)) % 360) * 60000))
    sx, sy = _spread_scale(el.w, s.spread), _spread_scale(el.h, s.spread)
    return (f'<a:effectLst><a:outerShdw blurRad="{emu(s.blur)}" dist="{emu(dist)}" dir="{direction}" sx="{sx}" sy="{sy}" algn="ctr" '
            f'rotWithShape="0">{_clr(s.color, alpha)}</a:outerShdw></a:effectLst>')


def _geometry_xml(el: Element) -> str:
    prst = {"rect": "rect", "roundRect": "roundRect", "ellipse": "ellipse", "round2SameRect": "round2SameRect"}.get(el.geometry, "rect")
    if el.kind == "ellipse":
        prst = "ellipse"
    gd = ""
    short = max(1e-6, min(el.w, el.h))
    if prst == "roundRect":
        gd = f'<a:gd name="adj" fmla="val {int(round(min(0.5, el.radius / short) * 100000))}"/>'
    elif prst == "round2SameRect":
        gd = f'<a:gd name="adj1" fmla="val {int(round(min(0.5, el.radius / short) * 100000))}"/><a:gd name="adj2" fmla="val 0"/>'
    return f'<a:prstGeom prst="{prst}"><a:avLst>{gd}</a:avLst></a:prstGeom>'


def _set_sppr(spPr, el: Element, fill: str | None) -> None:
    """Geometry, fill, outline and effects, in schema order."""
    for child in list(spPr):
        if child.tag != qn("a:xfrm"):
            spPr.remove(child)
    xml = _geometry_xml(el) + (fill or "") + line_xml(el) + effect_xml(el)
    for child in parse_xml(f"<a:wrap {nsdecls('a')}>{xml}</a:wrap>"):
        spPr.append(child)


def _name(shape, el: Element, descr: str = "") -> None:
    cNvPr = shape._element.xpath("./*[1]/p:cNvPr")[0]
    cNvPr.set("name", (el.name or el.kind)[:250])
    if descr:
        cNvPr.set("descr", descr[:250])


def _place(shape, el: Element) -> None:
    shape.left, shape.top, shape.width, shape.height = Emu(emu(el.x)), Emu(emu(el.y)), Emu(max(1, emu(el.w))), Emu(max(1, emu(el.h)))
    if el.rotation:
        shape.rotation = float(el.rotation) % 360


# ------------------------------------------------------------------ shapes
def add_shape(slide, el: Element):
    shape = slide.shapes.add_shape(_SHAPES.get("ellipse" if el.kind == "ellipse" else el.geometry, MSO_SHAPE.RECTANGLE), 0, 0, 1, 1)
    _place(shape, el)
    style = shape._element.find(qn("p:style"))
    if style is not None:  # theme line/fill/effect references would add outlines and shadows
        shape._element.remove(style)
    _set_sppr(shape._element.spPr, el, fill_xml(el.fill, el.opacity))
    _name(shape, el)
    return shape


def add_picture(slide, el: Element):
    pic = slide.shapes.add_picture(el.picture.path, Emu(emu(el.x)), Emu(emu(el.y)), Emu(max(1, emu(el.w))), Emu(max(1, emu(el.h))))
    if el.rotation:
        pic.rotation = float(el.rotation) % 360
    l, t, r, b = el.picture.crop
    if any(v > 1e-5 for v in (l, t, r, b)):
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = l, t, r, b
    if el.opacity < 0.999:
        blip = pic._element.blipFill.find(qn("a:blip"))
        blip.append(parse_xml(f'<a:alphaModFix {nsdecls("a")} amt="{int(round(el.opacity * 100000))}"/>'))
    _set_sppr(pic._element.spPr, el, None)
    _name(pic, el, el.name)
    return pic


# -------------------------------------------------------------------- text
def _rpr(t, tag: str = "a:rPr") -> str:
    size = max(100, min(400000, int(round(t.size * PT_PER_PX * 100))))
    spc = int(round(t.tracking * t.size * PT_PER_PX * 100))
    attrs = f'lang="en-US" sz="{size}" b="{1 if t.weight >= 600 else 0}" i="{1 if t.italic else 0}" dirty="0"'
    if spc:
        attrs += f' spc="{max(-400000, min(400000, spc))}"'
    if t.stroke:
        ln = f'<a:ln w="{emu(t.stroke_px)}"><a:solidFill>{_clr(t.stroke, t.alpha)}</a:solidFill></a:ln>'
        fill = "<a:noFill/>"
    else:
        ln, fill = "", f"<a:solidFill>{_clr(t.color, t.alpha)}</a:solidFill>"
    face = escape(t.family, {'"': "&quot;"})
    fonts = f'<a:latin typeface="{face}"/><a:ea typeface="{face}"/><a:cs typeface="{face}"/>'
    return f"<{tag} {attrs}>{ln}{fill}{fonts}</{tag}>"


def paragraphs(t) -> list[tuple[list[str], float]]:
    """Laid-out lines grouped into paragraphs: (lines, space before in px)."""
    if t.single_line:
        return [([" ".join(x for x in t.lines if x)], 0.0)]
    out: list[tuple[list[str], float]] = []
    current: list[str] = []
    gap = 0.0
    for line in t.lines:
        if line == "":
            if current:
                out.append((current, gap))
                current, gap = [], 0.0
            gap += t.gap_px
        else:
            current.append(line)
    if current:
        out.append((current, gap))
    return out


def text_body_xml(el: Element) -> str:
    t = el.text
    algn = {"left": "l", "center": "ctr", "right": "r"}.get(t.align, "l")
    line_pts = max(0, min(158400, int(round(t.line_px * PT_PER_PX * 100))))
    rpr = _rpr(t)
    paras = []
    for lines, gap in paragraphs(t):
        before = f'<a:spcBef><a:spcPts val="{int(round(gap * PT_PER_PX * 100))}"/></a:spcBef>' if gap else '<a:spcBef><a:spcPts val="0"/></a:spcBef>'
        ppr = f'<a:pPr algn="{algn}" marL="0" indent="0"><a:lnSpc><a:spcPts val="{line_pts}"/></a:lnSpc>{before}<a:spcAft><a:spcPts val="0"/></a:spcAft><a:buNone/></a:pPr>'
        runs = []
        if t.natural:  # one run per paragraph: wraps at the box width when edited
            runs.append(f"<a:r>{rpr}<a:t>{escape(' '.join(lines))}</a:t></a:r>")
        else:
            for i, line in enumerate(lines):
                if i:
                    runs.append(f"<a:br>{rpr}</a:br>")
                runs.append(f"<a:r>{rpr}<a:t>{escape(line)}</a:t></a:r>")
        paras.append(f"<a:p>{ppr}{''.join(runs)}{_rpr(t, 'a:endParaRPr')}</a:p>")
    wrap = "square" if t.natural else "none"
    body = (f'<a:bodyPr wrap="{wrap}" lIns="0" tIns="0" rIns="0" bIns="0" anchor="t" anchorCtr="0" rtlCol="0" vert="horz">'
            f"<a:noAutofit/></a:bodyPr><a:lstStyle/>")
    return f"<p:txBody {nsdecls('a', 'p')}>{body}{''.join(paras)}</p:txBody>"


def add_text(slide, el: Element):
    box = slide.shapes.add_textbox(Emu(emu(el.x)), Emu(emu(el.y)), Emu(max(1, emu(el.w))), Emu(max(1, emu(el.h))))
    if el.rotation:
        box.rotation = float(el.rotation) % 360
    sp = box._element
    old = sp.find(qn("p:txBody"))
    sp.replace(old, parse_xml(text_body_xml(el)))
    # a text box has no fill or outline of its own
    _set_sppr(sp.spPr, Element("rect", el.x, el.y, el.w, el.h), "<a:noFill/>")
    _name(box, el)
    return box


# -------------------------------------------------------------------- file
def write_pptx(paged: PagedScene, path: Path, title: str = "", notes: dict[int, str] | None = None) -> Path:
    """Write the paged scene; the file is re-opened and checked before it counts."""
    scene = paged.scene
    width_emu, height_emu = emu(scene.width), emu(paged.page_height)
    if width_emu > MAX_SLIDE_EMU or height_emu > MAX_SLIDE_EMU:
        raise ExportError("The page size is larger than PowerPoint allows (56 inches per side).")
    prs = PptxPresentation()
    prs.slide_width, prs.slide_height = Emu(width_emu), Emu(height_emu)
    blank = next((lay for lay in prs.slide_layouts if lay.name.lower() == "blank"), prs.slide_layouts[6])
    for page in paged.pages:
        slide = prs.slides.add_slide(blank)
        first = next((e for e in page.elements if e.role in ("background", "filler") and isinstance(e.fill, Solid)), None)
        if first is not None:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string(_hex(first.fill.color))
        for el in page.elements:
            if el.w <= 0.01 or el.h <= 0.01:
                continue
            if el.kind == "text" and el.text is not None:
                add_text(slide, el)
            elif el.kind == "image" and el.picture is not None:
                add_picture(slide, el)
            else:
                add_shape(slide, el)
        if notes and page.index in notes:
            slide.notes_slide.notes_text_frame.text = notes[page.index]
    props = prs.core_properties
    props.title = (title or scene.title or "Presentation")[:255]
    props.author = APP_NAME
    props.comments = f"Created with {APP_NAME} {__version__} for Canva. Page {scene.width} × {paged.page_height} px."
    path = Path(path)
    tmp = path.with_name(f".{path.name}.partial")
    prs.save(str(tmp))
    try:
        check = PptxPresentation(str(tmp))
        if len(check.slides) != len(paged.pages):
            raise ExportError("The PowerPoint file is incomplete.", f"slides={len(check.slides)} expected={len(paged.pages)}")
    except ExportError:
        tmp.unlink(missing_ok=True)
        raise
    except Exception as exc:  # noqa: BLE001
        tmp.unlink(missing_ok=True)
        raise ExportError("The PowerPoint file could not be verified.", str(exc)) from exc
    os.replace(tmp, path)
    return path
